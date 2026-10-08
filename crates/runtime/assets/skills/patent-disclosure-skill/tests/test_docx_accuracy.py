"""Regression checks for both independently distributed Word converters."""
from __future__ import annotations

import importlib.util
import io
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SCOPES = ("patent-disclosure", "patent-application")


def load(scope, name):
    path = ROOT / "skills" / scope / "tools" / (name + ".py")
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


class DocxAccuracyTests(unittest.TestCase):
    def each(self):
        for scope in SCOPES:
            math = load(scope, "math_to_omml")
            converter = load(scope, "md_to_docx")
            yield scope, math, converter

    def test_complete_control_words_and_glyph_identity(self):
        source = r"\left(x\right)\leq y\geq z\le a\ge b\land c\lor d"
        for scope, math, _ in self.each():
            with self.subTest(scope=scope):
                self.assertEqual(math.normalize_latex_for_omml(source),
                    r"\left(x\right)\leq y\geq z\leq a\geq b\wedge c\vee d")
                self.assertIn(r"\circ", math.normalize_latex_for_omml(r"x\circ y"))
                self.assertEqual(math.normalize_latex_for_omml(r"25^{\circ}C"), "25℃")
                superscript = math.latex_to_omml(r"a^{\circ}", display=False)
                self.assertIn("∘", "".join(superscript.itertext()))
                self.assertNotIn("°", "".join(superscript.itertext()))

    def test_native_math_structures(self):
        expressions = [
            (r"\left(\frac{x_1^2}{\sqrt{1+y}}\right)\leq 1", {"d": 1, "f": 1, "rad": 1, "sSubSup": 1}),
            (r"\hat{x}+\tilde{y}+\vec{z}+\bar{a}+\underline{b}", {"acc": 3, "bar": 2}),
            (r"\sum_{i=1}^n x_i", {"limLow": 1, "limUpp": 1}),
            (r"\lim_{n\to\infty} a_n", {"limLow": 1}),
            (r"\begin{pmatrix}a&b\\c&d\end{pmatrix}", {"d": 1, "m": 1, "mr": 2}),
            (r"{}^{n}", {"sSup": 1}),
        ]
        for scope, math, _ in self.each():
            for source, expected in expressions:
                with self.subTest(scope=scope, source=source):
                    result = math.latex_to_omml(source)
                    self.assertEqual(result.tag, qn("m:oMathPara"))
                    for tag, count in expected.items():
                        self.assertEqual(len(list(result.iter(qn("m:" + tag)))), count)
                    self.assertFalse(list(result.iter(qn("w:drawing"))))
            cases = math.latex_to_omml(r"\begin{cases}x&x\geq0\\-x&x<0\end{cases}")
            self.assertEqual(list(cases.iter(qn("m:begChr")))[0].get(qn("m:val")), "{")
            self.assertEqual(list(cases.iter(qn("m:endChr")))[0].get(qn("m:val")), "")

    def test_mathvariant_is_inherited(self):
        for scope, math, _ in self.each():
            with self.subTest(scope=scope):
                native = math.latex_to_omml(r"\mathrm{abc}+x", display=False)
                styles = {run.find(qn("m:t")).text: run.find(qn("m:rPr")).find(qn("m:sty")).get(qn("m:val"))
                          for run in native.iter(qn("m:r"))}
                for char in "abc":
                    self.assertEqual(styles[char], "p")
                self.assertEqual(styles["x"], "i")

    def test_unsupported_math_fails_instead_of_flattening(self):
        for scope, math, _ in self.each():
            for source in (r"\unknowncommand{x}", r"\overbrace{x+y}"):
                with self.subTest(scope=scope, source=source):
                    with self.assertRaises(ValueError):
                        math.latex_to_omml(source)
            target = math._element("oMath")
            with self.assertRaises(ValueError):
                math._append_mathml(target, ElementTree.fromstring(
                    '<menclose xmlns="http://www.w3.org/1998/Math/MathML"><mi>x</mi></menclose>'))

    def test_semantics_does_not_duplicate_annotation(self):
        for scope, math, _ in self.each():
            with self.subTest(scope=scope):
                node = ElementTree.fromstring(
                    '<semantics xmlns="http://www.w3.org/1998/Math/MathML"><mi>x</mi><annotation encoding="application/x-tex">x</annotation></semantics>')
                target = math._element("oMath")
                math._append_mathml(target, node)
                self.assertEqual("".join(target.itertext()), "x")

    def test_default_requires_editable_math_and_reports_explicit_fallback(self):
        for scope, _, converter in self.each():
            with self.subTest(scope=scope), redirect_stderr(io.StringIO()):
                with self.assertRaisesRegex(ValueError, "editable_math_required"):
                    converter.convert_md_to_docx(r"Bad $\unknowncommand{x}$.", None)
                converter.convert_md_to_docx(r"Bad $\unknowncommand{x}$.", None, require_editable_math=False)
                self.assertEqual(converter.get_math_stats().text, 1)
                self.assertEqual(converter.get_math_stats().omml, 0)
                with self.assertRaisesRegex(ValueError, "native math cannot be disabled"):
                    converter.convert_md_to_docx("text", None, prefer_omml=False, require_editable_math=True)

    def test_failed_conversion_and_bad_encoding_keep_existing_output(self):
        for scope, _, converter in self.each():
            with self.subTest(scope=scope), tempfile.TemporaryDirectory() as temporary:
                root = Path(temporary)
                source, output = root / "source.md", root / "result.docx"
                output.write_bytes(b"previous approved document")
                for content in (b'Bad $\\unknowncommand{x}$.', b'bad\xff'):
                    source.write_bytes(content)
                    with redirect_stderr(io.StringIO()), redirect_stdout(io.StringIO()):
                        self.assertEqual(converter.main(["-i", str(source), "-o", str(output)]), 1)
                    self.assertEqual(output.read_bytes(), b"previous approved document")

    def test_atomic_save_failure_and_native_package_roundtrip(self):
        for scope, math, converter in self.each():
            with self.subTest(scope=scope), tempfile.TemporaryDirectory() as temporary:
                output = Path(temporary) / "result.docx"
                output.write_bytes(b"previous")
                doc = Document()
                doc.add_paragraph()._p.append(math.latex_to_omml(r"\frac{a}{b}+\hat{x}"))
                with patch.object(type(doc), "save", side_effect=OSError("simulated full disk")):
                    with self.assertRaises(OSError):
                        converter.save_docx_atomically(doc, output)
                self.assertEqual(output.read_bytes(), b"previous")
                self.assertEqual(list(output.parent.iterdir()), [output])
                converter.save_docx_atomically(doc, output)
                reopened = Document(output)
                self.assertEqual(len(list(reopened.element.iter(qn("m:oMath")))), 1)
                with ZipFile(output) as archive:
                    self.assertIn(b"<m:acc>", archive.read("word/document.xml"))


if __name__ == "__main__":
    unittest.main()
